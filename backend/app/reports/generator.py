from __future__ import annotations

from pathlib import Path

from docx import Document
from sqlalchemy import text
from sqlalchemy.orm import Session


def build_report(db: Session, project_id: str, artifact_root: Path) -> Path:
    project = db.execute(text("SELECT name, description FROM projects WHERE id = :id"), {"id": project_id}).mappings().one()
    overlaps = db.execute(
        text(
            """
            SELECT source, theme, layer_name, area_ha, percent_aoi
            FROM overlaps
            WHERE project_id = :id
            ORDER BY theme, source, layer_name
            """
        ),
        {"id": project_id},
    ).mappings().all()
    report_dir = artifact_root / str(project_id) / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    path = report_dir / "relatorio_analise.docx"
    doc = Document()
    doc.add_heading(f"Relatório geoespacial — {project['name']}", level=1)
    if project["description"]:
        doc.add_paragraph(project["description"])
    doc.add_heading("Sobreposições", level=2)
    table = doc.add_table(rows=1, cols=5)
    for cell, header in zip(table.rows[0].cells, ["Fonte", "Tema", "Camada", "Área (ha)", "% AOI"], strict=True):
        cell.text = header
    for overlap in overlaps:
        row = table.add_row().cells
        row[0].text = overlap["source"]
        row[1].text = overlap["theme"]
        row[2].text = overlap["layer_name"]
        row[3].text = f"{overlap['area_ha']:.4f}"
        row[4].text = f"{overlap['percent_aoi']:.2f}"
    doc.save(path)
    return path
